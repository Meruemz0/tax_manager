"""Editable plain-text contract templates and portable DOCX/PDF renderers."""

from io import BytesIO
from pathlib import Path
import re
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


FIELDS = {
    "contract_number": "合同编号",
    "customer_name": "客户名称",
    "customer_tax_id": "客户税号",
    "customer_contact": "客户联系人",
    "customer_phone": "客户电话",
    "provider_name": "服务方名称",
    "provider_tax_id": "服务方税号",
    "provider_contact": "服务方联系人",
    "provider_phone": "服务方电话",
    "service_name": "服务项目",
    "order_date": "订单日期",
    "amount": "订单金额",
    "service_start_month": "服务开始月份",
    "service_end_month": "服务结束月份",
    "due_date": "应收日期",
    "sign_date": "签订日期",
    "materials_day": "每月资料提交日",
    "payment_terms": "付款约定",
    "extra_terms": "补充条款",
}
TOKEN = re.compile(r"{{\s*([a-z][a-z0-9_]{0,63})\s*}}")


def validate_template_body(body: str, field_keys=()) -> None:
    if not isinstance(body, str) or not body.strip() or len(body) > 20000:
        raise ValueError("模板正文不能为空，且最多 20000 字")
    unknown = set(TOKEN.findall(body)) - FIELDS.keys() - set(field_keys)
    if unknown:
        raise ValueError("未知占位符：" + "、".join(sorted(unknown)))
    remainder = TOKEN.sub("", body)
    if "{{" in remainder or "}}" in remainder:
        raise ValueError("占位符格式错误，请使用 {{customer_name}} 形式")


def render_tokens(body: str, values: dict[str, str]) -> str:
    validate_template_body(body)
    missing = set(TOKEN.findall(body)) - values.keys()
    if missing:
        raise ValueError("缺少占位符值：" + "、".join(sorted(missing)))
    return TOKEN.sub(lambda match: str(values[match.group(1)]), body)


def build_docx(title: str, body: str) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.6)
    section.right_margin = Cm(2.6)
    normal = document.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(11)
    normal.font._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.space_after = Pt(6)
    heading = document.add_heading(title, level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for paragraph in body.split("\n"):
        document.add_paragraph(paragraph or " ")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def build_pdf(title: str, body: str) -> bytes:
    font_name = "NotoSansSC"
    if font_name not in pdfmetrics.getRegisteredFontNames():
        font_file = Path(__file__).parent / "fonts" / "NotoSansSC.ttf"
        pdfmetrics.registerFont(TTFont(font_name, str(font_file)))
    coverage = pdfmetrics.getFont(font_name).face.charToGlyph
    unsupported = {char for char in title + body if char not in "\n\r\t" and (
        ord(char) > 0xFFFF or ord(char) not in coverage
    )}
    if unsupported:
        raise ValueError("PDF 字体无法完整显示部分字符，请下载 DOCX 核对并编辑")
    output = BytesIO()
    doc = SimpleDocTemplate(
        output, pagesize=A4, rightMargin=54, leftMargin=54,
        topMargin=50, bottomMargin=50,
        title=title, author="税务客户管理",
    )
    heading = ParagraphStyle(
        "contract-heading", fontName=font_name, fontSize=16,
        leading=23, alignment=TA_CENTER, spaceAfter=22, wordWrap="CJK",
    )
    text_style = ParagraphStyle(
        "contract-body", fontName=font_name, fontSize=10.5,
        leading=18, spaceAfter=6, wordWrap="CJK", textColor=colors.black,
    )
    story = [Paragraph(escape(title), heading)]
    for line in body.split("\n"):
        story.append(Paragraph(escape(line) if line else "&#160;", text_style))
    doc.build(story)
    return output.getvalue()
